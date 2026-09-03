import csv
import io
from django.db import transaction
from openpyxl import Workbook, load_workbook
from .models import Students
from academics.models import Program


EXPECTED_COLUMNS = [
    'first_name', 'last_name', 'email', 'phone',
    'date_of_birth', 'gender', 'address',
    'guardian_name', 'guardian_phone',
]

REQUIRED_COLUMNS = ['first_name', 'last_name', 'gender']


def generate_roll_number(program, batch):
    """
    Generate next sequential roll number for a program+batch combination.
    Format: {PROGRAM_CODE}-{BATCH}-{SEQUENCE}
    Example: BCA-2081-001, BCA-2081-002
    """
    prefix = f"{program.code}-{batch}"

    # Find the highest existing sequence number for this prefix
    existing = Students.objects.filter(
        roll_number__startswith=prefix
    ).order_by('-roll_number')

    if existing.exists():
        last_roll = existing.first().roll_number
        try:
            last_seq = int(last_roll.split('-')[-1])
        except (ValueError, IndexError):
            last_seq = 0
        next_seq = last_seq + 1
    else:
        next_seq = 1

    return f"{prefix}-{next_seq:03d}"


def parse_csv_file(file_obj):
    """
    Parse an uploaded CSV file and return a list of student dicts.
    Returns: (data_list, errors_list)
    """
    errors = []
    data = []

    try:
        text = file_obj.read().decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(text))

        # Normalize header names
        if reader.fieldnames:
            reader.fieldnames = [f.strip().lower().replace(' ', '_') for f in reader.fieldnames]

        # Check required columns
        missing = [col for col in REQUIRED_COLUMNS if col not in (reader.fieldnames or [])]
        if missing:
            errors.append(f"Missing required columns: {', '.join(missing)}")
            return data, errors

        for row_num, row in enumerate(reader, start=2):
            row_data, row_errors = _validate_row(row, row_num)
            if row_errors:
                errors.extend(row_errors)
            else:
                data.append(row_data)

    except UnicodeDecodeError:
        errors.append("File encoding error. Please save the CSV as UTF-8.")
    except Exception as e:
        errors.append(f"Error reading CSV: {str(e)}")

    return data, errors


def parse_excel_file(file_obj):
    """
    Parse an uploaded Excel (.xlsx) file and return a list of student dicts.
    Returns: (data_list, errors_list)
    """
    errors = []
    data = []

    try:
        wb = load_workbook(file_obj, read_only=True)
        ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            errors.append("The Excel file is empty.")
            return data, errors

        # Get headers from first row
        headers = [str(h).strip().lower().replace(' ', '_') if h else '' for h in rows[0]]

        # Check required columns
        missing = [col for col in REQUIRED_COLUMNS if col not in headers]
        if missing:
            errors.append(f"Missing required columns: {', '.join(missing)}")
            return data, errors

        for row_num, row in enumerate(rows[1:], start=2):
            row_dict = {}
            for i, header in enumerate(headers):
                if header and i < len(row):
                    val = row[i]
                    row_dict[header] = str(val).strip() if val is not None else ''
                elif header:
                    row_dict[header] = ''

            # Skip completely empty rows
            if all(v == '' for v in row_dict.values()):
                continue

            row_data, row_errors = _validate_row(row_dict, row_num)
            if row_errors:
                errors.extend(row_errors)
            else:
                data.append(row_data)

        wb.close()

    except Exception as e:
        errors.append(f"Error reading Excel file: {str(e)}")

    return data, errors


def _validate_row(row, row_num):
    """Validate a single row of student data."""
    errors = []
    data = {}

    # Required fields
    first_name = row.get('first_name', '').strip()
    last_name = row.get('last_name', '').strip()
    gender = row.get('gender', '').strip().upper()

    if not first_name:
        errors.append(f"Row {row_num}: first_name is required")
    if not last_name:
        errors.append(f"Row {row_num}: last_name is required")
    if gender not in ('M', 'F', 'O'):
        errors.append(f"Row {row_num}: gender must be M, F, or O (got '{gender}')")

    if errors:
        return data, errors

    data = {
        'first_name': first_name,
        'last_name': last_name,
        'gender': gender[0] if gender else 'M',
        'email': row.get('email', '').strip(),
        'phone': row.get('phone', '').strip(),
        'date_of_birth': row.get('date_of_birth', '').strip() or None,
        'address': row.get('address', '').strip(),
        'guardian_name': row.get('guardian_name', '').strip(),
        'guardian_phone': row.get('guardian_phone', '').strip(),
    }

    # Validate date format if provided
    if data['date_of_birth']:
        from datetime import datetime
        try:
            datetime.strptime(data['date_of_birth'], '%Y-%m-%d')
        except ValueError:
            try:
                # Try alternate format
                parsed = datetime.strptime(data['date_of_birth'], '%m/%d/%Y')
                data['date_of_birth'] = parsed.strftime('%Y-%m-%d')
            except ValueError:
                errors.append(
                    f"Row {row_num}: date_of_birth must be YYYY-MM-DD or MM/DD/YYYY format"
                )

    return data, errors


@transaction.atomic
def bulk_create_students(student_data_list, program, batch):
    """
    Create multiple students with auto-generated roll numbers.
    Returns: (created_students, errors)
    """
    created = []
    errors = []

    for i, data in enumerate(student_data_list):
        try:
            roll_number = generate_roll_number(program, batch)

            student = Students(
                roll_number=roll_number,
                first_name=data['first_name'],
                last_name=data['last_name'],
                email=data.get('email', ''),
                phone=data.get('phone', ''),
                date_of_birth=data.get('date_of_birth') or None,
                gender=data.get('gender', 'M'),
                address=data.get('address', ''),
                program=program,
                batch=batch,
                current_semester=1,
                guardian_name=data.get('guardian_name', ''),
                guardian_phone=data.get('guardian_phone', ''),
                status='active',
            )

            # Set default password
            default_password = student.set_default_password()
            student.save()

            created.append({
                'student': student,
                'roll_number': roll_number,
                'default_password': default_password,
            })

        except Exception as e:
            errors.append(f"Error creating student {data.get('first_name', '')} {data.get('last_name', '')}: {str(e)}")

    return created, errors


def generate_sample_csv():
    """Generate a sample CSV file content for download."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(EXPECTED_COLUMNS)
    writer.writerow([
        'Ram', 'Sharma', 'ram@example.com', '9841234567',
        '2005-03-15', 'M', 'Kathmandu, Nepal',
        'Hari Sharma', '9801234567',
    ])
    writer.writerow([
        'Sita', 'Gurung', 'sita@example.com', '9851234567',
        '2005-06-20', 'F', 'Pokhara, Nepal',
        'Gita Gurung', '9811234567',
    ])
    return output.getvalue()


def generate_sample_excel():
    """Generate a sample Excel file for download."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"

    # Header row
    ws.append(EXPECTED_COLUMNS)

    # Sample data
    ws.append([
        'Ram', 'Sharma', 'ram@example.com', '9841234567',
        '2005-03-15', 'M', 'Kathmandu, Nepal',
        'Hari Sharma', '9801234567',
    ])
    ws.append([
        'Sita', 'Gurung', 'sita@example.com', '9851234567',
        '2005-06-20', 'F', 'Pokhara, Nepal',
        'Gita Gurung', '9811234567',
    ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
